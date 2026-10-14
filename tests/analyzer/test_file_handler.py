# Copyright (c) Microsoft Corporation.
# Licensed under the MIT License.

"""Tests for file handler module in analyzer."""

import unittest
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import superbench.analyzer.file_handler as file_handler


@pytest.mark.parametrize('function', ['value', 'variance'])
@pytest.mark.parametrize('threshold', [-1.25, 0, 2.5])
@pytest.mark.parametrize(
    'operator,excel_operator',
    [('>', 'greaterThan'), ('<', 'lessThan'), ('>=', 'greaterThanOrEqual'),
     ('<=', 'lessThanOrEqual'), ('==', 'equal'), ('!=', 'notEqual')]
)
def test_excel_comparison_operators(function, threshold, operator, excel_operator):
    """Preserve all supported comparison operators in actual XLSX formatting."""
    rules = {'rule': {'function': function, 'metrics': {'bench/value': None},
                      'criteria': 'lambda x: x {} {}'.format(operator, threshold)}}
    with BytesIO() as output:
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            file_handler.output_excel_data_not_accept(writer, pd.DataFrame({'bench/value': [2.0]}), rules)
        output.seek(0)
        workbook = load_workbook(output)
        try:
            worksheet = workbook['Not Accept']
            conditions = [
                rule for cell_range in worksheet.conditional_formatting
                for rule in worksheet.conditional_formatting[cell_range] if rule.type == 'cellIs'
            ]
            assert len(conditions) == 1
            assert conditions[0].operator == excel_operator
            assert float(conditions[0].formula[0]) == threshold
        finally:
            workbook.close()


def test_excel_empty_results():
    """Keep empty diagnostic reports valid without comparison formatting."""
    rules = {'rule': {'function': 'value', 'metrics': {'bench/value': None}, 'criteria': 'lambda x: x >= 1'}}
    with BytesIO() as output:
        with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
            file_handler.output_excel_data_not_accept(writer, pd.DataFrame(columns=['bench/value']), rules)
        output.seek(0)
        workbook = load_workbook(output)
        try:
            assert not workbook['Not Accept'].conditional_formatting
        finally:
            workbook.close()


class TestFileHandler(unittest.TestCase):
    """Test for file handler."""
    def setUp(self):
        """Method called to prepare the test fixture."""
        self.parent_path = Path(__file__).parent
        self.test_rule_file_fake = str(self.parent_path / 'test_rules_fake.yaml')

    def tearDown(self):
        """Method called after the test method has been called and the result recorded."""
        for file in [self.test_rule_file_fake]:
            p = Path(file)
            if p.is_file():
                p.unlink()

    def test_file_handler(self):
        """Test for the file handler."""
        test_raw_data = str(self.parent_path / 'test_results.jsonl')
        test_rule_file = str(self.parent_path / 'test_rules.yaml')
        test_baseline_file = str(self.parent_path / 'test_baseline.json')
        test_raw_data_fake = str(self.parent_path / 'test_results_fake.jsonl')
        test_rule_file_fake = str(self.parent_path / 'test_rules_fake.yaml')
        test_aseline_file_fake = str(self.parent_path / 'test_baseline_fake.json')
        # Test - read_raw_data
        raw_data_df = file_handler.read_raw_data(test_raw_data)
        assert (not raw_data_df.empty)
        self.assertRaises(FileNotFoundError, file_handler.read_raw_data, test_raw_data_fake)
        # Test - read rules
        self.assertRaises(FileNotFoundError, file_handler.read_rules, test_rule_file_fake)
        rules = file_handler.read_rules(test_rule_file)
        assert (rules)
        # Test - read baseline
        self.assertRaises(FileNotFoundError, file_handler.read_baseline, test_aseline_file_fake)
        baseline = file_handler.read_baseline(test_baseline_file)
        assert (baseline)
        # Test - generate_md_table
        data_df = pd.DataFrame([[1, 2], [3, 4]])
        lines = file_handler.generate_md_table(data_df, header=['A', 'B'])
        expected_lines = ['| A | B |\n', '| --- | --- |\n', '| 1 | 2 |\n', '| 3 | 4 |\n']
        assert (lines == expected_lines)
